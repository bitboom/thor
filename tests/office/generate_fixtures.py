"""Create synthetic Office fixtures and modified copies in a caller-selected folder.

Requires the Codex bundled Python runtime with python-docx, openpyxl, python-pptx,
and Pillow. The outputs are test data, never real business documents.
"""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
from shutil import copyfile
from zipfile import ZipFile

from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches as DocInches, Pt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt as PptPt

CASES = [
    ('01-overview', '출시 준비 현황', [4, 7, 6]),
    ('02-operations', '운영 점검 기록', [3, 5, 8]),
    ('03-summary', '검토 결과 요약', [6, 2, 9]),
]
MARKER = '복사본 수정 확인'


def digest(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def image(path: Path) -> None:
    im = Image.new('RGB', (640, 220), '#f0f5ff')
    draw = ImageDraw.Draw(im)
    draw.rectangle((16, 16, 623, 203), outline='#234075', width=4)
    draw.polygon([(72, 165), (190, 55), (295, 145), (412, 66), (558, 166)], fill='#5b8bd9')
    draw.ellipse((492, 30, 554, 92), fill='#f2c65b')
    im.save(path, 'PNG')


def make_docx(path: Path, title: str, values: list[int], png: Path) -> None:
    d = Document()
    d.core_properties.author = 'Thor synthetic fixture'
    d.styles['Normal'].font.name = 'Malgun Gothic'
    d.styles['Normal'].font.size = Pt(11)
    d.add_heading(title, 0)
    d.add_paragraph('가상의 Lantern 프로젝트 데이터입니다. 원본을 보존하고 복사본만 수정합니다.')
    table = d.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    table.cell(0, 0).merge(table.cell(0, 1)).text = '구분 및 담당'
    table.cell(0, 2).text = '수량'
    for row, label, owner, value in zip(table.rows[1:], ['문서', '검토', '발표'], ['민아', '지우', '하늘'], values):
        row.cells[0].text, row.cells[1].text, row.cells[2].text = label, owner, str(value)
    d.add_paragraph('합성 이미지')
    d.add_picture(str(png), width=DocInches(4.5))
    d.save(path)


def make_xlsx(path: Path, title: str, values: list[int], png: Path) -> None:
    w = Workbook()
    w.properties.creator = 'Thor synthetic fixture'
    sh = w.active
    sh.title = '현황'
    sh.merge_cells('A1:C1')
    sh['A1'] = title
    sh['A1'].font = Font(name='Malgun Gothic', size=16, bold=True, color='183A68')
    sh['A1'].alignment = Alignment(horizontal='center')
    sh.append(['항목', '담당자', '수량'])
    for label, owner, value in zip(['문서', '검토', '발표'], ['민아', '지우', '하늘'], values):
        sh.append([label, owner, value])
    sh['A6'] = '합계'
    sh['C6'] = '=SUM(C3:C5)'
    for cell in sh[2]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='315C93')
    for col, width in [('A', 18), ('B', 18), ('C', 12)]:
        sh.column_dimensions[col].width = width
    chart = BarChart()
    chart.title = '항목별 수량'
    chart.add_data(Reference(sh, min_col=3, min_row=2, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(sh, min_col=1, min_row=3, max_row=5))
    sh.add_chart(chart, 'E2')
    picture = ExcelImage(str(png))
    picture.width, picture.height = 320, 110
    sh.add_image(picture, 'E18')
    w.save(path)


def make_pptx(path: Path, title: str, values: list[int], png: Path) -> None:
    prs = Presentation()
    prs.core_properties.author = 'Thor synthetic fixture'
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(11.8), Inches(0.7))
    run = box.text_frame.paragraphs[0].add_run()
    run.text, run.font.name, run.font.size, run.font.bold = title, 'Malgun Gothic', PptPt(28), True
    table = slide.shapes.add_table(4, 3, Inches(0.7), Inches(1.4), Inches(5.1), Inches(2.2)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text, table.cell(0, 2).text = '구분 및 담당', '수량'
    for row, label, owner, value in zip(range(1, 4), ['문서', '검토', '발표'], ['민아', '지우', '하늘'], values):
        table.cell(row, 0).text, table.cell(row, 1).text, table.cell(row, 2).text = label, owner, str(value)
    data = CategoryChartData()
    data.categories = ['문서', '검토', '발표']
    data.add_series('수량', values)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.5), Inches(1.3), Inches(5.8), Inches(3.5), data).chart
    # Normalize signed axis IDs into the schema's positive 31-bit range,
    # retaining the same mapping for every axis and cross-reference.
    axis_ids = {}
    for node in chart._chartSpace.xpath('.//c:axId | .//c:crossAx'):
        value = node.get('val')
        if value not in axis_ids:
            axis_ids[value] = str(100000001 + len(axis_ids))
        node.set('val', axis_ids[value])
    slide.shapes.add_picture(str(png), Inches(0.8), Inches(4.3), width=Inches(4.8))
    second = prs.slides.add_slide(prs.slide_layouts[6])
    text = second.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11), Inches(1.6))
    text.text_frame.text = '두 번째 슬라이드: 원본 유지와 복사본 수정 확인'
    text.text_frame.paragraphs[0].runs[0].font.size = PptPt(24)
    prs.save(path)


def build(out: Path) -> dict:
    if out.exists() and any(out.iterdir()):
        raise SystemExit(f'Output folder is not empty: {out}')
    out.mkdir(parents=True, exist_ok=True)
    originals, copies = out / 'originals', out / 'copies'
    originals.mkdir(); copies.mkdir()
    png = out / 'synthetic-image.png'
    image(png)
    entries = []
    for stem, title, values in CASES:
        for ext, builder in [('docx', make_docx), ('xlsx', make_xlsx), ('pptx', make_pptx)]:
            source = originals / f'{stem}.{ext}'
            target = copies / source.name
            builder(source, title, values, png)
            before = digest(source)
            copyfile(source, target)
            if ext == 'docx':
                d = Document(target); d.add_paragraph(MARKER); d.save(target)
            elif ext == 'xlsx':
                w = load_workbook(target); w['현황']['A8'] = MARKER; w.save(target)
            else:
                p = Presentation(target); p.slides[1].shapes.add_textbox(Inches(0.8), Inches(3), Inches(8), Inches(1)).text_frame.text = MARKER; p.save(target)
            assert digest(source) == before and digest(target) != before
            entries.append({'file': source.name, 'title': title, 'values': values, 'originalSha256': before, 'copySha256': digest(target)})
    manifest = {'synthetic': True, 'copyMarker': MARKER, 'entries': entries}
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    verify(out)
    return manifest


def verify(out: Path) -> None:
    manifest = json.loads((out / 'manifest.json').read_text(encoding='utf-8'))
    assert len(manifest['entries']) == 9
    for item in manifest['entries']:
        source, copy = out / 'originals' / item['file'], out / 'copies' / item['file']
        assert digest(source) == item['originalSha256'] and digest(copy) == item['copySha256']
        assert digest(source) != digest(copy)
        ext = source.suffix
        with ZipFile(copy) as package:
            names = package.namelist()
            assert any('/media/' in name for name in names), item['file']
            if ext in ('.xlsx', '.pptx'):
                assert any('/charts/chart' in name and name.endswith('.xml') for name in names), item['file']
        if ext == '.docx':
            d = Document(copy); assert MARKER in '\n'.join(p.text for p in d.paragraphs); assert len(d.tables) == 1
        elif ext == '.xlsx':
            w = load_workbook(copy); sh = w['현황']; assert sh['A8'].value == MARKER and sh['C6'].value == '=SUM(C3:C5)'; assert 'A1:C1' in str(sh.merged_cells)
        else:
            p = Presentation(copy); assert len(p.slides) == 2; assert any(MARKER in sh.text for sh in p.slides[1].shapes if sh.has_text_frame)
    print(f'Verified {len(manifest["entries"])} synthetic originals and modified copies; original hashes unchanged.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output-dir', required=True, type=Path)
    parser.add_argument('--verify-only', action='store_true')
    args = parser.parse_args()
    verify(args.output_dir) if args.verify_only else build(args.output_dir)
